"""Text bodies inside shapes → SVG <text>/<tspan> elements.

We do our own line breaking against measured advance widths (Pillow over
the real TTFs) instead of trusting the SVG renderer to wrap — typst's
resvg pipeline does not wrap text at all.  Every run gets an explicit x
so metric differences between Pillow and the final renderer cannot
accumulate across a line.
"""

from __future__ import annotations

import xml.sax.saxutils as _sax
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from typstpresenter.convert.pptx_style import shape_font_rgb

_EMU_PER_PT = 12700.0
_DEFAULT_SIZE_PT = 18.0  # PowerPoint default for text in autoshapes
_DEFAULT_FONT = "Calibri"

# default text insets (EMU): l/r 91440, t/b 45720
_DEF_INS = (91440 / _EMU_PER_PT, 45720 / _EMU_PER_PT)

_FONT_FILES = {
    ("calibri", False, False): "calibri.ttf",
    ("calibri", True, False): "calibrib.ttf",
    ("calibri", False, True): "calibrii.ttf",
    ("calibri", True, True): "calibriz.ttf",
    ("arial", False, False): "arial.ttf",
    ("arial", True, False): "arialbd.ttf",
    ("arial", False, True): "ariali.ttf",
    ("arial", True, True): "arialbi.ttf",
}


@lru_cache(maxsize=64)
def _load_font(name: str, size_pt: float, bold: bool, italic: bool):
    from PIL import ImageFont

    fname = _FONT_FILES.get((name.lower(), bold, italic)) or _FONT_FILES.get(
        (name.lower(), False, False), "calibri.ttf"
    )
    fonts_dir = Path("C:/Windows/Fonts")
    try:
        # 4x supersampling for sub-pt advance precision
        return ImageFont.truetype(str(fonts_dir / fname), round(size_pt * 4))
    except OSError:
        try:
            return ImageFont.truetype(str(fonts_dir / "calibri.ttf"), round(size_pt * 4))
        except OSError:
            return ImageFont.load_default()


def text_width_pt(text: str, font_name: str, size_pt: float, bold: bool, italic: bool) -> float:
    font = _load_font(font_name, size_pt, bold, italic)
    try:
        return font.getlength(text) / 4.0
    except AttributeError:
        return len(text) * size_pt * 0.5


def font_line_metrics(font_name: str, size_pt: float, bold: bool, italic: bool) -> tuple[float, float]:
    """(ascent, descent) in pt."""
    font = _load_font(font_name, size_pt, bold, italic)
    try:
        ascent, descent = font.getmetrics()
        return ascent / 4.0, descent / 4.0
    except AttributeError:
        return size_pt * 0.8, size_pt * 0.25


@dataclass
class Run:
    text: str
    font: str
    size_pt: float
    bold: bool
    italic: bool
    color: str  # "#RRGGBB"
    underline: bool = False

    def width(self) -> float:
        return text_width_pt(self.text, self.font, self.size_pt, self.bold, self.italic)


@dataclass
class Line:
    runs: list[Run] = field(default_factory=list)
    align: int | None = None  # PP_ALIGN or None

    def width(self) -> float:
        return sum(r.width() for r in self.runs)

    def ascent(self) -> float:
        return max((font_line_metrics(r.font, r.size_pt, r.bold, r.italic)[0] for r in self.runs),
                   default=0.0)

    def pitch(self) -> float:
        return max(
            (sum(font_line_metrics(r.font, r.size_pt, r.bold, r.italic)) for r in self.runs),
            default=0.0,
        )


_BREAK = None  # sentinel in run lists: explicit a:br line break


def _resolve_runs(paragraph, shape, default_color: str,
                  default_size: float, size_scale: float) -> list[Run | None]:
    """Paragraph content in document order: Run items, _BREAK for a:br."""
    from pptx.oxml.ns import qn

    from typstpresenter.convert.pptx_inherit import resolve_font_size_pt
    from typstpresenter.diagram2svg.symbols import (
        is_symbol_font,
        map_symbol_text,
        run_symbol_font,
    )

    runs: list[Run | None] = []
    by_el = {run._r: run for run in paragraph.runs}
    for child in paragraph._p:
        if child.tag == qn("a:br"):
            runs.append(_BREAK)
            continue
        if child.tag == qn("a:fld"):
            # field runs (slide numbers, dates) are invisible to
            # python-pptx's .runs; emit their cached text, default style
            fld_text = "".join(t.text or "" for t in child.findall(qn("a:t")))
            if fld_text:
                runs.append(Run(fld_text, _DEFAULT_FONT,
                                default_size * size_scale,
                                False, False, default_color))
            continue
        run = by_el.get(child)
        if run is None or not run.text:
            continue
        font = run.font
        size = resolve_font_size_pt(run, paragraph, shape)
        if size is None:
            size = default_size
        size *= size_scale
        color = None
        try:
            if font.color and font.color.type is not None and font.color.rgb is not None:
                color = f"#{font.color.rgb}"
        except (AttributeError, ValueError):
            pass
        name = font.name
        text = run.text
        sym = run_symbol_font(run)
        if is_symbol_font(name) or is_symbol_font(sym) or any(
                0xF000 <= ord(c) <= 0xF0FF for c in text):
            text = map_symbol_text(text, name if is_symbol_font(name) else sym)
            name = None  # mapped to Unicode; use the default font
        if not name or name.startswith("+"):  # +mn-lt / +mj-lt theme refs
            name = _DEFAULT_FONT
        runs.append(
            Run(
                text=text,
                font=name,
                size_pt=size,
                bold=bool(font.bold),
                italic=bool(font.italic),
                color=color or default_color,
                underline=bool(font.underline),
            )
        )
    return runs


def _wrap(runs: list[Run], max_w: float) -> list[list[Run]]:
    """Greedy word wrap preserving run styling."""
    # explode into word tokens (keeping trailing spaces attached)
    tokens: list[Run] = []
    for r in runs:
        parts = r.text.split(" ")
        for i, part in enumerate(parts):
            word = part + (" " if i < len(parts) - 1 else "")
            if word:
                tokens.append(Run(word, r.font, r.size_pt, r.bold, r.italic, r.color))
    lines: list[list[Run]] = []
    cur: list[Run] = []
    cur_w = 0.0
    for tok in tokens:
        w = tok.width()
        if cur and cur_w + text_width_pt(tok.text.rstrip(), tok.font, tok.size_pt,
                                         tok.bold, tok.italic) > max_w:
            lines.append(cur)
            cur, cur_w = [], 0.0
        cur.append(tok)
        cur_w += w
    if cur:
        lines.append(cur)
    # merge adjacent tokens with identical style back into single runs
    merged: list[list[Run]] = []
    for line in lines:
        out: list[Run] = []
        for tok in line:
            if out and (out[-1].font, out[-1].size_pt, out[-1].bold,
                        out[-1].italic, out[-1].color) == (
                    tok.font, tok.size_pt, tok.bold, tok.italic, tok.color):
                out[-1].text += tok.text
            else:
                out.append(Run(tok.text, tok.font, tok.size_pt, tok.bold, tok.italic, tok.color))
        merged.append(out)
    return merged


def layout_shape_text(
    shape,
    w_pt: float,
    h_pt: float,
    text_rect_pt: tuple[float, float, float, float] | None = None,
    default_size: float = _DEFAULT_SIZE_PT,
) -> list[str]:
    """SVG <text> elements (shape-local coordinates) for a shape's text.

    `text_rect_pt` is the preset geometry's text rectangle (l, t, r, b) in
    shape-local pt — PowerPoint lays text out inside it, not the full bbox.
    `default_size` is the deck-resolved size for runs whose inheritance
    chain yields nothing (real decks; generated decks set sizes explicitly).
    """
    if not getattr(shape, "has_text_frame", False):
        return []
    tf = shape.text_frame
    if not tf.text.strip():
        return []

    from typstpresenter.convert.pptx_inherit import resolve_alignment, resolve_bullet
    from typstpresenter.convert.textbody import autofit_scales

    size_scale, _ = autofit_scales(shape)

    rl, rt, rr, rb = text_rect_pt or (0.0, 0.0, w_pt, h_pt)
    default_color = f"#{shape_font_rgb(shape) or '000000'}"
    ins_x, ins_y = _DEF_INS
    box_w = max(rr - rl - 2 * ins_x, 1.0)

    # OOXML default anchor is top (bodyPr without @anchor)
    anchor = tf.vertical_anchor or MSO_ANCHOR.TOP

    lines: list[Line] = []
    for para in tf.paragraphs:
        try:
            align = resolve_alignment(para, shape)
        except (AttributeError, KeyError, ValueError):
            align = para.alignment
        content = _resolve_runs(para, shape, default_color, default_size, size_scale)
        real = [r for r in content if r is not None]
        if not real:
            lines.append(Line([], align))
            continue
        try:
            bullet = resolve_bullet(para, shape)
        except (AttributeError, KeyError, ValueError):
            bullet = None
        if bullet:
            from typstpresenter.diagram2svg.symbols import bullet_font, map_symbol_text

            glyph = map_symbol_text(bullet, bullet_font(para))
            r0 = real[0]
            content.insert(0, Run(f"{glyph} ", r0.font, r0.size_pt, False, False, r0.color))
        # explicit a:br breaks split the paragraph into wrap segments
        segment: list[Run] = []
        segments: list[list[Run]] = []
        for item in content:
            if item is _BREAK:
                segments.append(segment)
                segment = []
            else:
                segment.append(item)
        segments.append(segment)
        for seg in segments:
            if not seg:
                lines.append(Line([], align))
                continue
            for line_runs in _wrap(seg, box_w):
                lines.append(Line(line_runs, align))

    total_h = sum(
        ln.pitch() if ln.runs else _DEFAULT_SIZE_PT * 1.22 for ln in lines
    )
    if anchor == MSO_ANCHOR.MIDDLE:
        y = rt + (rb - rt - total_h) / 2.0
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = rb - ins_y - total_h
    else:
        y = rt + ins_y

    elements: list[str] = []
    for ln in lines:
        pitch = ln.pitch() if ln.runs else _DEFAULT_SIZE_PT * 1.22
        if ln.runs:
            baseline = y + ln.ascent()
            # OOXML default paragraph alignment is left
            align = ln.align if ln.align is not None else PP_ALIGN.LEFT
            lw = ln.width()
            if align == PP_ALIGN.RIGHT:
                x = rr - ins_x - lw
            elif align in (PP_ALIGN.CENTER, PP_ALIGN.JUSTIFY):
                x = rl + (rr - rl - lw) / 2.0
            else:
                x = rl + ins_x
            spans = []
            for r in ln.runs:
                style = f'font-family="{r.font}" font-size="{r.size_pt:g}" fill="{r.color}"'
                if r.bold:
                    style += ' font-weight="bold"'
                if r.italic:
                    style += ' font-style="italic"'
                if r.underline:
                    style += ' text-decoration="underline"'
                spans.append(
                    f'<tspan x="{x:.3f}" y="{baseline:.3f}" {style}>'
                    f"{_sax.escape(r.text)}</tspan>"
                )
                x += r.width()
            elements.append("<text>" + "".join(spans) + "</text>")
        y += pitch
    return elements
